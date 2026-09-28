import argparse
import json
import math
import sqlite3
import struct
import unicodedata

from openpyxl import load_workbook


def normalize(value):
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return " ".join("".join(ch if ch.isalnum() else " " for ch in text.upper()).split())


def zone_for(zone, locality):
    if zone:
        return str(zone).strip()
    loc = normalize(locality)
    if loc == "KENNEDY":
        return "ZONA OCCIDENTE"
    if loc.startswith("RAFAEL URIBE"):
        return "ZONA SUR"
    return "SIN ZONA"


def fix_lat(value):
    value = float(value)
    return abs(value) / 10_000_000 if abs(value) > 90 else value


def fix_lon(value):
    value = float(value)
    return -abs(value) / 10_000_000 if abs(value) > 180 else value


def read_attended(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["PARQUES"]
    rows = ws.iter_rows(values_only=True)
    headers = [normalize(v) for v in next(rows)]
    index = {name: i for i, name in enumerate(headers)}
    result = []
    for row in rows:
        try:
            name = row[index["PARQUE"]]
            locality = row[index["LOCALIDAD"]]
            zone = row[index["ZONA"]]
            lat = fix_lat(row[index["LATITUD"]])
            lon = fix_lon(row[index["LONGITUD"]])
        except (KeyError, TypeError, ValueError):
            continue
        result.append({
            "name": str(name).strip(),
            "key": normalize(name),
            "locality": "" if locality is None else str(locality).strip(),
            "zone": zone_for(zone, locality),
            "lat": lat,
            "lon": lon,
        })
    return result


def gpkg_wkb(blob):
    flags = blob[3]
    envelope = (flags >> 1) & 7
    envelope_size = {0: 0, 1: 32, 2: 48, 3: 48, 4: 64}.get(envelope, 0)
    return memoryview(blob)[8 + envelope_size:]


def read_uint(data, offset, endian):
    return struct.unpack_from(endian + "I", data, offset)[0], offset + 4


def read_point(data, offset, endian):
    x, y = struct.unpack_from(endian + "dd", data, offset)
    return [x, y], offset + 16


def parse_geometry(data, offset=0):
    endian = "<" if data[offset] == 1 else ">"
    offset += 1
    raw_type, offset = read_uint(data, offset, endian)
    geom_type = raw_type & 0xFF if raw_type & 0xE0000000 else raw_type % 1000
    if geom_type == 3:
        ring_count, offset = read_uint(data, offset, endian)
        rings = []
        for _ in range(ring_count):
            point_count, offset = read_uint(data, offset, endian)
            ring = []
            for _ in range(point_count):
                point, offset = read_point(data, offset, endian)
                ring.append(point)
            rings.append(ring)
        return {"type": "Polygon", "coordinates": rings}, offset
    if geom_type == 6:
        polygon_count, offset = read_uint(data, offset, endian)
        polygons = []
        for _ in range(polygon_count):
            polygon, offset = parse_geometry(data, offset)
            polygons.append(polygon["coordinates"])
        return {"type": "MultiPolygon", "coordinates": polygons}, offset
    raise ValueError(f"Unsupported geometry type {raw_type}")


def point_segment_distance(point, start, end):
    px, py = point
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    if dx == 0 and dy == 0:
        return math.hypot(px - x1, py - y1)
    t = max(0.0, min(1.0, ((px - x1) * dx + (py - y1) * dy) / (dx * dx + dy * dy)))
    return math.hypot(px - (x1 + t * dx), py - (y1 + t * dy))


def simplify_line(points, tolerance):
    if len(points) <= 4:
        return points
    closed = points[0] == points[-1]
    work = points[:-1] if closed else points
    if len(work) <= 3:
        return points

    def rdp(seq):
        if len(seq) <= 2:
            return seq
        best_i, best_d = 0, 0.0
        for i in range(1, len(seq) - 1):
            distance = point_segment_distance(seq[i], seq[0], seq[-1])
            if distance > best_d:
                best_i, best_d = i, distance
        if best_d <= tolerance:
            return [seq[0], seq[-1]]
        return rdp(seq[: best_i + 1])[:-1] + rdp(seq[best_i:])

    reduced = rdp(work + [work[0]])
    if reduced[0] != reduced[-1]:
        reduced.append(reduced[0])
    return reduced if len(reduced) >= 4 else points


def simplify_geometry(geometry, tolerance):
    if geometry["type"] == "Polygon":
        coordinates = [[list(map(lambda v: round(v, 6), p)) for p in simplify_line(ring, tolerance)] for ring in geometry["coordinates"]]
    else:
        coordinates = [
            [[list(map(lambda v: round(v, 6), p)) for p in simplify_line(ring, tolerance)] for ring in polygon]
            for polygon in geometry["coordinates"]
        ]
    return {"type": geometry["type"], "coordinates": coordinates}


def bbox(geometry):
    if geometry["type"] == "Polygon":
        points = [p for ring in geometry["coordinates"] for p in ring]
    else:
        points = [p for polygon in geometry["coordinates"] for ring in polygon for p in ring]
    xs = [p[0] for p in points]
    ys = [p[1] for p in points]
    return min(xs), min(ys), max(xs), max(ys)


def point_in_ring(x, y, ring):
    inside = False
    j = len(ring) - 1
    for i in range(len(ring)):
        xi, yi = ring[i]
        xj, yj = ring[j]
        if (yi > y) != (yj > y):
            crossing = (xj - xi) * (y - yi) / (yj - yi) + xi
            if x < crossing:
                inside = not inside
        j = i
    return inside


def contains(geometry, x, y):
    polygons = [geometry["coordinates"]] if geometry["type"] == "Polygon" else geometry["coordinates"]
    return any(polygon and point_in_ring(x, y, polygon[0]) for polygon in polygons)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--gpkg", required=True)
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--tolerance", type=float, default=0.00008)
    args = parser.parse_args()

    attended = read_attended(args.xlsx)
    by_name = {item["key"]: item for item in attended}
    connection = sqlite3.connect(f"file:{args.gpkg}?mode=ro", uri=True)
    rows = connection.execute(
        'select "geom", "Nombre_del_Parque", "Nombre_Localidad", "Tipo_de_Parque", "Codigo_Parque" from parques'
    )
    parsed = []
    for geom_blob, name, locality, park_type, code in rows:
        geometry, _ = parse_geometry(gpkg_wkb(geom_blob))
        parsed.append({
            "geometry": geometry,
            "bbox": bbox(geometry),
            "name": "" if name is None else str(name).strip(),
            "key": normalize(name),
            "locality": "" if locality is None else str(locality).strip(),
            "park_type": "" if park_type is None else str(park_type).strip(),
            "code": code,
        })

    selected = {}
    for item in attended:
        matches = [feature for feature in parsed if feature["key"] == item["key"]]
        if not matches:
            matches = []
            for feature in parsed:
                min_x, min_y, max_x, max_y = feature["bbox"]
                if min_x <= item["lon"] <= max_x and min_y <= item["lat"] <= max_y and contains(feature["geometry"], item["lon"], item["lat"]):
                    matches.append(feature)
        for feature in matches:
            selected[(feature["code"], feature["name"])] = (feature, item)

    features = []
    for feature, item in selected.values():
        features.append({
            "type": "Feature",
            "properties": {
                "PARQUE": item["name"],
                "ZONA": item["zone"],
                "LOCALIDAD": item["locality"],
                "TIPO": feature["park_type"],
                "CODIGO": feature["code"],
            },
            "geometry": simplify_geometry(feature["geometry"], args.tolerance),
        })

    payload = {"type": "FeatureCollection", "features": features}
    with open(args.output, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))
    print(json.dumps({"attended_points": len(attended), "matched_polygons": len(features), "output": args.output}, ensure_ascii=False))


if __name__ == "__main__":
    main()
